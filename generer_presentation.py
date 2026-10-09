"""Genere une presentation PowerPoint a partir des rapports du projet GTFS."""

from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
REPORTS = ROOT / "reports"
NAVY = RGBColor(22, 48, 74)
TEAL = RGBColor(23, 107, 135)
ORANGE = RGBColor(217, 121, 65)
INK = RGBColor(42, 54, 66)
MUTED = RGBColor(98, 112, 124)
PALE = RGBColor(239, 245, 248)
WHITE = RGBColor(255, 255, 255)
SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)


def read_json(path: Path) -> dict[str, Any]:
    if not path.is_file():
        raise FileNotFoundError(f"Rapport introuvable : {path}")
    with path.open(encoding="utf-8") as source:
        return json.load(source)


def read_csv(path: Path) -> list[dict[str, str]]:
    if not path.is_file():
        raise FileNotFoundError(f"Tableau introuvable : {path}")
    with path.open(encoding="utf-8-sig", newline="") as source:
        return list(csv.DictReader(source))


def format_number(value: int | float, decimals: int = 0) -> str:
    formatted = f"{value:,.{decimals}f}"
    return formatted.replace(",", " ").replace(".", ",")


def format_date(value: str) -> str:
    if len(value) == 8 and value.isdigit():
        return f"{value[6:8]}/{value[4:6]}/{value[:4]}"
    return value


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    width: float,
    height: float,
    *,
    size: int = 18,
    color: RGBColor = INK,
    bold: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
) -> None:
    box = slide.shapes.add_textbox(
        Inches(x), Inches(y), Inches(width), Inches(height)
    )
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.04)
    frame.margin_right = Inches(0.04)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.text = text
    paragraph.font.name = "Aptos"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color


def add_bullets(
    slide,
    items: list[str],
    x: float,
    y: float,
    width: float,
    height: float,
    *,
    size: int = 19,
    color: RGBColor = INK,
) -> None:
    box = slide.shapes.add_textbox(
        Inches(x), Inches(y), Inches(width), Inches(height)
    )
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.08)
    frame.margin_right = Inches(0.08)
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = f"•  {item}"
        paragraph.level = 0
        paragraph.space_after = Pt(12)
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = color


def add_title(slide, title: str, subtitle: str | None = None) -> None:
    add_text(slide, title, 0.7, 0.35, 11.9, 0.55, size=27, color=NAVY, bold=True)
    if subtitle:
        add_text(slide, subtitle, 0.72, 0.98, 11.7, 0.42, size=13, color=MUTED)
    rule = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.72), Inches(1.45), Inches(0.8), Inches(0.06)
    )
    rule.fill.solid()
    rule.fill.fore_color.rgb = TEAL
    rule.line.fill.background()


def add_footer(slide, slide_number: int) -> None:
    add_text(
        slide,
        "Analyse de l’offre de transport collectif • Pays de la Loire",
        0.72,
        7.12,
        10.7,
        0.2,
        size=9,
        color=MUTED,
    )
    add_text(slide, str(slide_number), 12.0, 7.08, 0.55, 0.25, size=10, color=MUTED, align=PP_ALIGN.RIGHT)


def new_content_slide(prs: Presentation, title: str, subtitle: str | None = None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, title, subtitle)
    add_footer(slide, len(prs.slides))
    return slide


def add_metric_card(slide, x: float, y: float, value: str, label: str, accent=TEAL) -> None:
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(2.7),
        Inches(1.25),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = PALE
    shape.line.fill.background()
    add_text(slide, value, x + 0.15, y + 0.16, 2.4, 0.45, size=25, color=accent, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, label, x + 0.15, y + 0.68, 2.4, 0.38, size=12, color=MUTED, align=PP_ALIGN.CENTER)


def add_picture(slide, filename: str, x: float, y: float, width: float, height: float) -> None:
    path = REPORTS / filename
    if not path.is_file():
        raise FileNotFoundError(f"Graphique introuvable : {path}")
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(width), height=Inches(height))


def add_table(slide, headers: list[str], rows: list[list[str]], x: float, y: float, width: float, height: float) -> None:
    table = slide.shapes.add_table(
        len(rows) + 1,
        len(headers),
        Inches(x),
        Inches(y),
        Inches(width),
        Inches(height),
    ).table
    table.first_row = True
    for column, header in enumerate(headers):
        cell = table.cell(0, column)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.bold = True
            paragraph.font.size = Pt(12)
            paragraph.font.color.rgb = WHITE
    for row_index, row in enumerate(rows, start=1):
        for column, value in enumerate(row):
            cell = table.cell(row_index, column)
            cell.text = value
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if row_index % 2 else PALE
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.name = "Aptos"
                paragraph.font.size = Pt(11)
                paragraph.font.color.rgb = INK


def build_presentation(output: Path) -> None:
    offer = read_json(REPORTS / "analyse_offre.json")
    model = read_json(REPORTS / "etape6_modeles_risques.json")
    h1 = read_csv(REPORTS / "etape4_h1_offre_par_periode.csv")
    h2 = read_csv(REPORTS / "etape4_h2_offre_jours_ouvres_weekend.csv")
    h3 = read_csv(REPORTS / "etape4_h3_correlation_lignes_departs.csv")

    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = NAVY
    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.18), SLIDE_HEIGHT
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = TEAL
    accent.line.fill.background()
    add_text(slide, "MOBILITÉ • PAYS DE LA LOIRE", 0.85, 0.85, 10.7, 0.45, size=15, color=RGBColor(153, 214, 222), bold=True)
    add_text(slide, "Analyser l’offre\npour éclairer la desserte", 0.85, 1.65, 11.4, 1.8, size=36, color=WHITE, bold=True)
    add_text(
        slide,
        "Analyse exploratoire et évaluation de prévisions à partir des données GTFS",
        0.9,
        4.0,
        10.9,
        0.8,
        size=20,
        color=WHITE,
    )
    add_text(
        slide,
        f"Période étudiée : {format_date(offer['date_min'])} – {format_date(offer['date_max'])}",
        0.9,
        6.55,
        10.5,
        0.35,
        size=13,
        color=RGBColor(203, 216, 226),
    )

    slide = new_content_slide(prs, "Question métier et périmètre")
    add_text(slide, "Comment repérer les arrêts et créneaux où l’offre théorique est la plus faible, puis estimer le niveau d’offre attendu ?", 0.85, 1.85, 7.0, 1.2, size=23, color=NAVY, bold=True)
    add_bullets(
        slide,
        [
            "Indicateur : nombre de départs planifiés par arrêt, date et période horaire.",
            "Source principale : données GTFS du réseau DESTINEO.",
            "Les résultats décrivent l’offre planifiée, pas la fréquentation, la demande ni les retards réellement observés.",
        ],
        0.9,
        3.35,
        7.0,
        2.4,
        size=17,
    )
    add_metric_card(slide, 9.0, 2.05, format_number(offer["observation_count"]), "observations arrêt-date-période")
    add_metric_card(slide, 9.0, 3.65, format_number(offer["date_count"]), "dates de service", accent=ORANGE)

    slide = new_content_slide(prs, "Un socle de données contrôlé")
    add_text(slide, "Chaîne de traitement", 0.9, 1.8, 5.7, 0.4, size=20, color=NAVY, bold=True)
    add_bullets(
        slide,
        [
            "Fichiers GTFS : lignes, trajets, arrêts et horaires.",
            "Contrôle et nettoyage, puis structuration relationnelle dans SQLite.",
            "Vues analytiques pour agréger les départs par arrêt et période.",
            "Exploration, baseline temporelle et comparaison de modèles.",
        ],
        0.9,
        2.35,
        6.3,
        3.2,
        size=16,
    )
    add_text(slide, "Base de référence", 7.7, 1.8, 4.4, 0.4, size=20, color=NAVY, bold=True)
    add_metric_card(slide, 7.7, 2.4, "4,0 M", "passages horaires contrôlés")
    add_metric_card(slide, 10.7, 2.4, "20 170", "arrêts")
    add_metric_card(slide, 7.7, 4.0, "182 725", "trajets planifiés", accent=ORANGE)
    add_metric_card(slide, 10.7, 4.0, "1 029", "lignes")
    add_text(slide, f"Couverture analytique : {format_date(offer['date_min'])} au {format_date(offer['date_max'])}.", 7.75, 5.8, 4.8, 0.55, size=13, color=MUTED)

    slide = new_content_slide(prs, "L’offre varie selon l’heure et le type de jour")
    add_picture(slide, "etape4_h1_offre_par_periode.png", 0.75, 1.7, 6.0, 3.35)
    add_picture(slide, "etape4_h2_offre_jours_ouvres_weekend.png", 6.85, 1.7, 5.7, 3.35)
    h1_sorted = sorted(h1, key=lambda row: float(row["average_departures"]), reverse=True)
    top_period = h1_sorted[0]
    low_period = h1_sorted[-1]
    weekday = next(row for row in h2 if row["day_type"] == "Jour ouvré")
    weekend = next(row for row in h2 if row["day_type"] == "Week-end")
    decrease = (1 - float(weekend["average_departures"]) / float(weekday["average_departures"])) * 100
    add_text(
        slide,
        f"Pointes les plus élevées : {top_period['time_period']} ({format_number(float(top_period['average_departures']), 2)} départs moyens). "
        f"Creux : {low_period['time_period']} ({format_number(float(low_period['average_departures']), 2)}).",
        0.85,
        5.28,
        11.7,
        0.65,
        size=15,
        color=INK,
    )
    add_text(
        slide,
        f"Jours ouvrés : {format_number(float(weekday['average_departures']), 2)} départs moyens ; week-end : "
        f"{format_number(float(weekend['average_departures']), 2)}, soit une baisse d’environ {format_number(decrease, 1)} %.",
        0.85,
        6.0,
        11.7,
        0.5,
        size=15,
        color=TEAL,
        bold=True,
    )

    slide = new_content_slide(prs, "Le nombre de lignes est associé à l’offre")
    add_picture(slide, "etape4_h3_correlation_lignes_departs.png", 0.8, 1.75, 6.2, 3.65)
    correlation = h3[0]
    add_metric_card(slide, 7.65, 2.0, format_number(float(correlation["pearson_correlation"]), 4), "corrélation de Pearson")
    add_metric_card(slide, 7.65, 3.55, format_number(float(correlation["spearman_correlation"]), 4), "corrélation de Spearman", accent=ORANGE)
    add_text(slide, "La relation positive est modérée : plusieurs lignes vont souvent de pair avec davantage de départs, sans démontrer une causalité ni expliquer seules le niveau d’offre.", 7.65, 5.2, 4.5, 1.1, size=16, color=INK)

    slide = new_content_slide(prs, "Prédire : la baseline historique reste la référence")
    models = {item["name"]: item for item in model["models"]}
    rows = [
        ["Baseline historique", format_number(model["baseline_cv"]["mean_mae"], 4), format_number(model["baseline_holdout"]["mae"], 4), format_number(model["baseline_holdout"]["rmse"], 4)],
        ["Random Forest", format_number(models["random_forest"]["cv_mae"], 4), format_number(models["random_forest"]["test_mae"], 4), format_number(models["random_forest"]["test_rmse"], 4)],
        ["Hist. Gradient Boosting", format_number(models["hist_gradient_boosting"]["cv_mae"], 4), format_number(models["hist_gradient_boosting"]["test_mae"], 4), format_number(models["hist_gradient_boosting"]["test_rmse"], 4)],
    ]
    add_table(slide, ["Candidat", "MAE validation\n(3 fenêtres)", "MAE test final", "RMSE test final"], rows, 0.9, 1.9, 11.5, 2.15)
    add_text(slide, f"Choix selon la validation chronologique : {model['selected_model']} (MAE moyenne {format_number(model['selected_cv_mae'], 4)}).", 1.0, 4.55, 11.1, 0.55, size=17, color=TEAL, bold=True)
    add_text(
        slide,
        f"Évaluation finale réservée : {format_number(model['test_rows'])} observations, "
        f"{format_number(model['test_date_count'])} dates, du {format_date(model['temporal_evaluation']['holdout_date_min'])} "
        f"au {format_date(model['temporal_evaluation']['holdout_date_max'])}. "
        "La MAE mesure l’écart absolu moyen en départs planifiés ; plus elle est faible, mieux c’est.",
        1.0,
        5.25,
        11.1,
        0.9,
        size=15,
        color=INK,
    )
    add_text(slide, "Le choix du modèle est fondé sur la validation, et non sur l’optimisation a posteriori du test final.", 1.0, 6.3, 11.1, 0.45, size=13, color=MUTED)

    slide = new_content_slide(prs, "Limites et pistes pour la décision")
    risks = model["risks"]
    risk_items = [
        f"{risk['category'].replace('_', ' ').capitalize()} : {risk['risk']}"
        for risk in risks
    ]
    add_bullets(slide, risk_items, 0.9, 1.85, 11.5, 3.8, size=15)
    add_text(
        slide,
        "À retenir : utiliser ces résultats comme aide à l’analyse, puis les compléter par la fréquentation réelle, les retards, la demande locale et une validation métier avant tout arbitrage de desserte.",
        0.95,
        5.85,
        11.3,
        0.85,
        size=17,
        color=TEAL,
        bold=True,
    )

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    print(f"Présentation créée : {output}")
    print(f"Nombre de diapositives : {len(prs.slides)}")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Génère une présentation PowerPoint à partir des rapports GTFS."
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=REPORTS / "presentation_analyse_mobilite.pptx",
        help="Chemin du PowerPoint à créer (défaut : reports/presentation_analyse_mobilite.pptx).",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    output = args.output if args.output.is_absolute() else ROOT / args.output
    build_presentation(output)


if __name__ == "__main__":
    main()
